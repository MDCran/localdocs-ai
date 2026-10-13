from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from localdocs.ingestion import (
    DocumentError,
    Page,
    chunk_pages,
    extract,
    safe_filename,
    validate_file,
)
from reportlab.pdfgen import canvas


def test_pdf_preserves_page_numbers(tmp_path: Path):
    path = tmp_path / "guide.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 700, "The archive retention period is 30 days.")
    pdf.showPage()
    pdf.drawString(72, 700, "Encryption keys rotate every 90 days.")
    pdf.save()
    pages = extract(path, ".pdf")
    assert [p.number for p in pages] == [1, 2]
    assert "30 days" in pages[0].text
    assert "90 days" in pages[1].text


def test_scanned_pdf_has_actionable_error(tmp_path: Path):
    path = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.rect(20, 20, 100, 100)
    pdf.save()
    with pytest.raises(DocumentError, match="requires OCR"):
        extract(path, ".pdf")


def test_docx_includes_headings_and_tables(tmp_path: Path):
    doc = Document()
    doc.add_heading("Release policy", level=1)
    doc.add_paragraph("Ship on Fridays.")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Operations"
    path = tmp_path / "policy.docx"
    doc.save(path)
    text = extract(path, ".docx")[0].text
    assert "# Release policy" in text and "Owner | Operations" in text


@pytest.mark.parametrize("extension", [".txt", ".md"])
def test_utf8_text_and_markdown(tmp_path: Path, extension: str):
    path = tmp_path / f"note{extension}"
    path.write_text("# Café\n\nResearch notes.", encoding="utf-8-sig")
    assert "Café" in extract(path, extension)[0].text


@pytest.mark.parametrize(
    "extension,data",
    [
        (".pdf", b"not a pdf"),
        (".docx", b"not a zip"),
        (".txt", b"binary\x00data"),
        (".md", b"\xff"),
        (".exe", b"MZ"),
        (".txt", b""),
    ],
)
def test_validation_rejects_disguised_files(tmp_path: Path, extension: str, data: bytes):
    path = tmp_path / "file"
    path.write_bytes(data)
    with pytest.raises(DocumentError):
        validate_file(path, extension)


def test_docx_archive_expansion_limit(tmp_path: Path):
    path = tmp_path / "bomb.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", b"a" * 50_000_001)
    with pytest.raises(DocumentError, match="safe parsing limit"):
        validate_file(path, ".docx")


def test_chunks_preserve_page_boundaries_and_section_titles():
    pages = [
        Page("# Storage\n\n" + "Keep local archives. " * 30, 3),
        Page("# Access\n\nRead only.", 4),
    ]
    chunks = chunk_pages(pages, 220, 40)
    assert len(chunks) > 2
    assert chunks[0].text.startswith("# Storage")
    assert all(c.heading == "Storage" for c in chunks if c.page == 3)
    assert chunks[-1].page == 4 and chunks[-1].heading == "Access"
    assert all(len(c.text) <= 220 for c in chunks)


def test_path_traversal_and_control_characters_are_removed():
    assert safe_filename("../../secret.pdf") == "secret.pdf"
    assert safe_filename("C:\\fakepath\\report.pdf") == "report.pdf"
    assert "\n" not in safe_filename("note\n.txt")
