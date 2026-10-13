import json
import re
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document as DocxDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from localdocs.config import get_settings


class DocumentError(ValueError):
    pass


@dataclass(frozen=True)
class Page:
    text: str
    number: int | None


@dataclass(frozen=True)
class TextChunk:
    text: str
    page: int | None
    heading: str | None


def extract_isolated(path: Path, extension: str) -> list[Page]:
    try:
        result = subprocess.run(
            [sys.executable, "-m", "localdocs.parser", str(path.resolve()), extension],
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=60,
            check=False,
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentError("Document parsing exceeded the 60-second limit") from exc
    try:
        payload = json.loads(result.stdout)
    except ValueError as exc:
        raise DocumentError("Document parser stopped unexpectedly; try a smaller document") from exc
    if result.returncode or "error" in payload:
        raise DocumentError(payload.get("error", "Document parsing failed"))
    return [Page(**page) for page in payload["pages"]]


def safe_filename(name: str) -> str:
    name = name.replace("\\", "/").split("/")[-1]
    return re.sub(r"[^\w. ()-]", "_", name)[:240] or "document"


def validate_file(path: Path, extension: str) -> None:
    if extension not in {".pdf", ".docx", ".txt", ".md"}:
        raise DocumentError("Supported formats are PDF, DOCX, TXT, and Markdown")
    with path.open("rb") as handle:
        prefix = handle.read(1024)
    if not prefix:
        raise DocumentError("The uploaded file is empty")
    if extension == ".pdf" and not prefix.startswith(b"%PDF-"):
        raise DocumentError("The file is not a valid PDF")
    if extension == ".docx":
        try:
            with ZipFile(path) as archive:
                entries = archive.infolist()
                if sum(e.file_size for e in entries) > 50_000_000 or len(entries) > 5000:
                    raise DocumentError("DOCX archive expands beyond the safe parsing limit")
                if "word/document.xml" not in archive.namelist():
                    raise DocumentError("The file is not a Word document")
                if any("vbaProject" in e.filename for e in entries):
                    raise DocumentError("Macro-enabled documents are not supported")
        except BadZipFile as exc:
            raise DocumentError("The Word document is corrupt") from exc
    if extension in {".txt", ".md"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeError as exc:
            raise DocumentError("Text documents must use UTF-8 encoding") from exc
        if "\x00" in text or any(ord(c) < 9 for c in text):
            raise DocumentError("The file contains binary content")


def clean(text: str) -> str:
    return re.sub(r"[ \t]+", " ", text.replace("\x00", "").replace("\r\n", "\n")).strip()


def extract(path: Path, extension: str) -> list[Page]:
    validate_file(path, extension)
    try:
        if extension == ".pdf":
            reader = PdfReader(path, strict=True)
            if reader.is_encrypted:
                raise DocumentError("Password-protected PDFs are not supported")
            if len(reader.pages) > 2000:
                raise DocumentError("PDF exceeds the 2,000-page limit")
            pages = [
                Page(clean(page.extract_text() or ""), n + 1) for n, page in enumerate(reader.pages)
            ]
        elif extension == ".docx":
            doc = DocxDocument(str(path))
            lines = [
                ("# " if p.style and p.style.name.startswith("Heading") else "") + p.text
                for p in doc.paragraphs
            ]
            lines.extend(
                " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
            )
            pages = [Page(clean("\n\n".join(lines)), None)]
        else:
            pages = [Page(clean(path.read_text(encoding="utf-8-sig")), None)]
    except DocumentError:
        raise
    except Exception as exc:
        raise DocumentError("Document parsing failed; the file may be corrupt") from exc
    if sum(len(p.text) for p in pages) > get_settings().max_extracted_chars:
        raise DocumentError("Extracted text exceeds the processing limit")
    if not any(p.text for p in pages):
        message = (
            "No extractable text. This PDF may be scanned and requires OCR."
            if extension == ".pdf"
            else "Document contains no extractable text"
        )
        raise DocumentError(message)
    return pages


def chunk_pages(pages: list[Page], size: int, overlap: int) -> list[TextChunk]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=size,
        chunk_overlap=overlap,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks: list[TextChunk] = []
    for page in pages:
        sections = re.split(r"(?m)(?=^#{1,6} )", page.text)
        for section in sections:
            heading = section.splitlines()[0].lstrip("# ") if section.startswith("#") else None
            for text in splitter.split_text(section):
                if text.strip():
                    chunks.append(TextChunk(text, page.number, heading))
    return chunks
